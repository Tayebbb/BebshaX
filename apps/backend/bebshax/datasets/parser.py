"""Safe dataset parser supporting CSV, TSV, JSON, JSONL, and structured tables."""

from __future__ import annotations

import csv
import io
import json
import lzma
import math
import re
import struct
import zipfile
import zlib
from typing import Any
from xml.etree.ElementTree import ParseError
from xml.parsers import expat

from bebshax.datasets.security import MAX_DATASET_FILE_SIZE_BYTES

# Shape caps. The 25 MB byte ceiling bounds the *input*, not what parsing
# expands it into: a tiny file can still describe millions of rows or thousands
# of columns, and every downstream profiler/segmenter pass is O(rows x columns).
# Exceeding a cap is an explicit error, never a silent truncation.
MAX_DATASET_ROWS = 200_000
MAX_DATASET_COLUMNS = 512
MAX_DATASET_CELLS = 1_000_000
MAX_DATASET_NESTING = 16
MAX_DATASET_NODES = 2_000_000
MAX_DATASET_NORMALIZED_BYTES = 50 * 1024 * 1024

# XLSX is a zip archive, so a few hundred KB can declare gigabytes of sheet XML.
MAX_XLSX_UNCOMPRESSED_BYTES = 200 * 1024 * 1024
MAX_XLSX_COMPRESSION_RATIO = 200
MAX_XLSX_ARCHIVE_MEMBERS = 1024
MAX_XLSX_XML_TOKEN_BYTES = 64 * 1024
_CELL_REFERENCE = re.compile(r"^\$?([A-Z]{1,3})\$?([1-9][0-9]{0,7})$")
_ZIP_END = struct.Struct("<4s4H2IH")
_ZIP_MEMBER = struct.Struct("<4s6H3I5H2I")


class DatasetParseError(Exception):
    """Raised when dataset content cannot be parsed."""

    def __init__(
        self, detail: str, *, error_code: str = "dataset_unparseable", status_code: int = 422,
    ) -> None:
        super().__init__(detail)
        self.detail = detail
        self.error_code = error_code
        self.status_code = status_code


def _limit(detail: str) -> DatasetParseError:
    return DatasetParseError(detail, error_code="dataset_limits_exceeded", status_code=413)


def _preflight_text_bytes(content: bytes) -> None:
    """Bound escaped-text expansion without allocating the decoded document."""
    size = 0
    for character in content:
        if character in (34, 92):
            size += 2
        elif character < 32 and character not in (9, 10, 13):
            size += 6
        elif character < 128:
            size += 1
        elif character >= 240:
            size += 12
        elif character >= 192:
            size += 6
        if size > MAX_DATASET_NORMALIZED_BYTES:
            raise _limit("Dataset exceeds the normalized byte budget.")


def _enforce_shape_counts(column_count: int, row_count: int) -> None:
    if column_count > MAX_DATASET_COLUMNS:
        raise _limit(f"Dataset exceeds the {MAX_DATASET_COLUMNS}-column limit.")
    if row_count > MAX_DATASET_ROWS:
        raise _limit(f"Dataset has more than {MAX_DATASET_ROWS:,} rows. Split it before uploading.")
    if column_count * row_count > MAX_DATASET_CELLS:
        raise _limit(f"Dataset exceeds the {MAX_DATASET_CELLS:,} aggregate cell limit.")


def _enforce_shape_caps(columns: list[str], rows: list[dict[str, Any]]) -> None:
    _enforce_shape_counts(len(columns), len(rows))


def _json_string_size(value: str) -> int:
    size = 2
    for character in value:
        codepoint = ord(character)
        if character in ('"', "\\", "\b", "\f", "\n", "\r", "\t"):
            size += 2
        elif codepoint < 32 or 127 <= codepoint <= 65535:
            size += 6
        elif codepoint > 65535:
            size += 12
        else:
            size += 1
        if size > MAX_DATASET_NORMALIZED_BYTES:
            raise _limit("Dataset exceeds the normalized byte budget.")
    return size


class _NormalizedBudget:
    """Upper bound for default json.dump output, including repeated keys/nulls."""

    def __init__(self, columns: list[str], row_count: int) -> None:
        _enforce_shape_counts(len(columns), row_count)
        self.column_count = len(columns)
        self.nodes = 0
        row_overhead = 2 + sum(_json_string_size(column) + 2 for column in columns)
        row_overhead += 2 * max(0, len(columns) - 1)
        self.size = 2 + row_count * row_overhead + 2 * max(0, row_count - 1)
        self._check()

    def _check(self) -> None:
        if self.size > MAX_DATASET_NORMALIZED_BYTES:
            raise _limit("Dataset exceeds the normalized byte budget.")

    def _value_size(self, value: Any, depth: int = 0) -> int:
        self.nodes += 1
        if self.nodes > MAX_DATASET_NODES:
            raise _limit("Dataset exceeds the nested complexity budget.")
        if depth > MAX_DATASET_NESTING:
            raise _limit("Dataset exceeds the nesting limit.")
        if value is None:
            return 4
        if isinstance(value, bool):
            return 4 if value else 5
        if isinstance(value, (int, float)):
            if isinstance(value, float) and not math.isfinite(value):
                raise DatasetParseError("Dataset numeric values must be finite.")
            return len(str(value))
        if isinstance(value, dict):
            size = 2 + 2 * max(0, len(value) - 1)
            for key, nested in value.items():
                size += _json_string_size(key) + 2 + self._value_size(nested, depth + 1)
                if size > MAX_DATASET_NORMALIZED_BYTES:
                    raise _limit("Dataset exceeds the normalized byte budget.")
            return size
        if isinstance(value, list):
            size = 2 + 2 * max(0, len(value) - 1)
            for nested in value:
                size += self._value_size(nested, depth + 1)
                if size > MAX_DATASET_NORMALIZED_BYTES:
                    raise _limit("Dataset exceeds the normalized byte budget.")
            return size
        return max(4, _json_string_size(value if isinstance(value, str) else str(value)))

    def add_record(self, record: dict[str, Any]) -> None:
        self.size += 4 * (self.column_count - len(record))
        for value in record.values():
            self.size += self._value_size(value)
            self._check()
        self._check()


def _preflight_json(content: bytes) -> None:
    stack: list[list[int]] = []
    in_string = False
    escaped = False
    scalar = False
    nodes = 0
    roots = 0
    for character in memoryview(content)[3 if content.startswith(b"\xef\xbb\xbf") else 0:]:
        if in_string:
            if escaped:
                escaped = False
            elif character == 92:
                escaped = True
            elif character == 34:
                in_string = False
            continue
        if character in b" \t\r\n,:":
            scalar = False
            if character == 58 and stack and stack[-1][0] == 123:
                stack[-1][1] += 1
                _enforce_shape_counts(stack[-1][1], 0)
            continue
        if character in (125, 93):
            expected = 123 if character == 125 else 91
            if not stack or stack.pop()[0] != expected:
                raise DatasetParseError("Invalid JSON dataset.")
            scalar = False
            continue
        if character in (123, 91, 34) or not scalar:
            nodes += 1
            if nodes > MAX_DATASET_NODES:
                raise _limit("Dataset exceeds the nested complexity budget.")
            if not stack:
                roots += 1
                _enforce_shape_counts(0, roots)
            elif stack[-1][0] == 91:
                stack[-1][1] += 1
                _enforce_shape_counts(0, stack[-1][1])
        if character in (123, 91):
            stack.append([character, 0])
            if len(stack) > MAX_DATASET_NESTING:
                raise _limit("Dataset exceeds the nesting limit.")
            scalar = False
        elif character == 34:
            in_string = True
            scalar = False
        else:
            scalar = True
    if stack or in_string:
        raise DatasetParseError("Invalid JSON dataset.")


def _unique_object(pairs: list[tuple[str, Any]]) -> dict[str, Any]:
    record: dict[str, Any] = {}
    for key, value in pairs:
        if key in record:
            raise DatasetParseError("JSON object fields must be unique.")
        record[key] = value
    return record


def _finite_float(value: str) -> float:
    number = float(value)
    if not math.isfinite(number):
        raise DatasetParseError("Dataset numeric values must be finite.")
    return number


def _reject_constant(value: str) -> None:
    raise DatasetParseError("Dataset numeric values must be finite.")


def _load_json(text: str) -> Any:
    try:
        return json.loads(
            text, object_pairs_hook=_unique_object,
            parse_float=_finite_float, parse_constant=_reject_constant,
        )
    except ValueError:
        raise DatasetParseError("Invalid JSON dataset.") from None


def _normalize_columns(columns: list[str]) -> list[str]:
    normalized = [str(column).strip() for column in columns]
    if not normalized or any(not column for column in normalized):
        raise DatasetParseError("Dataset column names must not be empty.")
    if len(set(normalized)) != len(normalized):
        raise DatasetParseError("Dataset column names must be unique after trimming whitespace.")
    _enforce_shape_caps(normalized, [])
    return normalized


def _normalize_record(record: dict[str, Any]) -> dict[str, Any]:
    columns = _normalize_columns(list(record))
    return dict(zip(columns, record.values(), strict=True))


def detect_format(content: bytes, filename: str = "", content_type: str = "") -> str:
    """Detect dataset format from file name, content-type header, or magic bytes."""
    fname_lower = filename.lower()
    ctype_lower = content_type.lower()

    if fname_lower.endswith(".json") or "application/json" in ctype_lower:
        return "json"
    if fname_lower.endswith(".jsonl") or "application/x-ndjson" in ctype_lower:
        return "jsonl"
    if fname_lower.endswith(".tsv") or "text/tab-separated-values" in ctype_lower:
        return "tsv"
    if fname_lower.endswith(".xlsx") or fname_lower.endswith(".xls") or "spreadsheet" in ctype_lower:
        return "xlsx"
    if fname_lower.endswith(".csv") or "text/csv" in ctype_lower:
        return "csv"

    # Inspect first few non-empty characters
    snippet = content[:1024].decode("utf-8", errors="ignore").strip()
    if snippet.startswith("[") or (snippet.startswith("{") and "\n" not in snippet):
        return "json"
    if snippet.startswith("{") and "\n" in snippet:
        return "jsonl"
    if "\t" in snippet.splitlines()[0] if snippet.splitlines() else False:
        return "tsv"
    return "csv"


def parse_dataset_bytes(
    content: bytes,
    file_type: str | None = None,
    filename: str = "",
    content_type: str = "",
) -> tuple[list[str], list[dict[str, Any]]]:
    """Parse raw dataset bytes into column names and list of row dicts.
    
    Returns:
        (columns: list[str], rows: list[dict[str, Any]])
    """
    if not content:
        raise DatasetParseError("Dataset content is empty.")
    if len(content) > MAX_DATASET_FILE_SIZE_BYTES:
        raise _limit("Dataset input exceeds the file byte limit.")

    detected_type = file_type or detect_format(content, filename=filename, content_type=content_type)
    if detected_type != "xlsx":
        _preflight_text_bytes(content)

    try:
        if detected_type == "json":
            columns, rows = _parse_json(content)
        elif detected_type == "jsonl":
            columns, rows = _parse_jsonl(content)
        elif detected_type == "tsv":
            columns, rows = _parse_delimited(content, delimiter="\t")
        elif detected_type == "xlsx":
            columns, rows = _parse_xlsx(content)
        else:
            columns, rows = _parse_delimited(content, delimiter=",")
    except (UnicodeDecodeError, csv.Error):
        raise DatasetParseError("Dataset encoding or delimited format is invalid.") from None

    _enforce_shape_caps(columns, rows)
    return columns, rows


def _parse_delimited(content: bytes, delimiter: str = ",") -> tuple[list[str], list[dict[str, Any]]]:
    stream = io.TextIOWrapper(io.BytesIO(content), encoding="utf-8-sig", newline="")
    first_line = next((line for line in stream if line.strip()), "")
    if not first_line:
        raise DatasetParseError("Delimited text dataset contains no data rows.")

    # Check for automatic delimiter detection if ambiguous
    if delimiter == "," and "\t" in first_line and first_line.count("\t") > first_line.count(","):
        delimiter = "\t"
    elif delimiter == "," and ";" in first_line and first_line.count(";") > first_line.count(","):
        delimiter = ";"
    elif delimiter == "," and "|" in first_line and first_line.count("|") > first_line.count(","):
        delimiter = "|"

    row_count = _preflight_delimited(content, delimiter)
    stream.seek(0)
    reader = csv.reader(stream, delimiter=delimiter, strict=True)
    fieldnames = next((row for row in reader if row), None)
    if not fieldnames:
        raise DatasetParseError("Dataset header is missing or invalid.")

    clean_columns = _normalize_columns(fieldnames)
    budget = _NormalizedBudget(clean_columns, row_count)

    rows: list[dict[str, Any]] = []
    for row in reader:
        if not row:
            continue
        if len(row) > len(clean_columns):
            raise DatasetParseError("Dataset row contains more fields than its header.")
        record = {
            column: row[index] if index < len(row) else None
            for index, column in enumerate(clean_columns)
        }
        budget.add_record(record)
        _enforce_shape_counts(len(clean_columns), len(rows) + 1)
        rows.append({column: _cast_value(value) for column, value in record.items()})

    if not rows:
        raise DatasetParseError("Dataset contains headers but 0 data rows.")

    return clean_columns, rows


def _preflight_delimited(content: bytes, delimiter: str) -> int:
    separator = ord(delimiter)
    columns = 1
    header_columns = 0
    records = 0
    in_quotes = False
    closed_quote = False
    field_start = True
    record_started = False
    for character in memoryview(content)[3 if content.startswith(b"\xef\xbb\xbf") else 0:]:
        if in_quotes:
            if character == 34:
                in_quotes = False
                closed_quote = True
            continue
        if character == 34 and (field_start or closed_quote):
            in_quotes = True
            closed_quote = False
            field_start = False
            record_started = True
            continue
        closed_quote = False
        if character in (10, 13):
            if record_started:
                records += 1
                header_columns = header_columns or columns
                _enforce_shape_counts(header_columns, records - 1)
            columns = 1
            field_start = True
            record_started = False
        elif character == separator:
            columns += 1
            _enforce_shape_counts(columns, 0)
            field_start = True
            record_started = True
        else:
            field_start = False
            record_started = True
    if record_started:
        records += 1
        header_columns = header_columns or columns
    _enforce_shape_counts(header_columns, max(0, records - 1))
    return max(0, records - 1)


def _normalize_records(records: list[Any]) -> tuple[list[str], list[dict[str, Any]]]:
    _enforce_shape_counts(0, len(records))
    columns_seen: dict[str, None] = {}
    for index, item in enumerate(records):
        if not isinstance(item, dict):
            raise DatasetParseError("Every dataset row must be a JSON object.")
        record = _normalize_record(item)
        records[index] = record
        columns_seen.update(dict.fromkeys(record))
        _enforce_shape_counts(len(columns_seen), len(records))
    columns = list(columns_seen)
    if not columns:
        raise DatasetParseError("JSON objects have no valid keys/fields.")
    budget = _NormalizedBudget(columns, len(records))
    for record in records:
        budget.add_record(record)
    for index, record in enumerate(records):
        records[index] = {column: _cast_value(record.get(column)) for column in columns}
    return columns, records


def _parse_json(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    _preflight_json(content)
    data = _load_json(content.decode("utf-8-sig"))

    if isinstance(data, dict):
        # Look for common container keys: "data", "records", "rows", "items", "results"
        for key in ("data", "records", "rows", "items", "results", "personas", "users", "observations"):
            if key in data and isinstance(data[key], list):
                data = data[key]
                break
        else:
            data = [data]

    if not isinstance(data, list):
        raise DatasetParseError("JSON dataset must be an array of records/objects.")

    if not data:
        raise DatasetParseError("JSON dataset is an empty array.")

    return _normalize_records(data)


def _parse_jsonl(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    _preflight_json(content)
    records: list[dict[str, Any]] = []
    stream = io.TextIOWrapper(io.BytesIO(content), encoding="utf-8-sig")
    for line in stream:
        if not line.strip():
            continue
        _enforce_shape_counts(0, len(records) + 1)
        item = _load_json(line)
        if not isinstance(item, dict):
            raise DatasetParseError("Every dataset row must be a JSON object.")
        records.append(item)
    if not records:
        raise DatasetParseError("JSONL dataset contains no valid JSON object lines.")
    return _normalize_records(records)


class _SpreadsheetPreflight:
    """SAX counters bound XML, sparse coordinates and shared-string expansion."""

    def __init__(self) -> None:
        self.nodes = 0
        self.text_bytes = 0
        self.string_sizes: list[int] = []
        self.references: dict[int, int] = {}
        self.tags: list[str] = []
        self.sheet_rows = 0
        self.sheet_columns = 0
        self.current_row = 0
        self.current_column = 0
        self.shared_size = 0
        self.shared_cell = False
        self.reference = ""

    def _coordinate(self, value: str) -> None:
        match = _CELL_REFERENCE.fullmatch(value)
        if match is None:
            raise DatasetParseError("Spreadsheet contains invalid cell coordinates.")
        column = 0
        for character in match[1]:
            column = column * 26 + ord(character) - ord("A") + 1
        row = int(match[2])
        self.sheet_rows = max(self.sheet_rows, row)
        self.sheet_columns = max(self.sheet_columns, column)
        self.current_column = column
        _enforce_shape_counts(self.sheet_columns, max(0, self.sheet_rows - 1))

    def start(self, name: str, attributes: dict[str, str]) -> None:
        tag = name.rsplit("}", 1)[-1]
        self.tags.append(tag)
        self.nodes += 1
        if self.nodes > MAX_DATASET_NODES:
            raise _limit("Spreadsheet exceeds the nested complexity budget.")
        if len(self.tags) > MAX_DATASET_NESTING:
            raise _limit("Spreadsheet exceeds the nesting limit.")
        if tag == "worksheet":
            self.sheet_rows = self.sheet_columns = self.current_row = self.current_column = 0
        elif tag == "dimension" and "worksheet" in self.tags:
            for coordinate in attributes.get("ref", "").split(":"):
                self._coordinate(coordinate)
        elif tag == "row" and "worksheet" in self.tags:
            row = attributes.get("r", str(self.current_row + 1))
            if not row.isascii() or not row.isdigit() or len(row) > 8 or int(row) < 1:
                raise DatasetParseError("Spreadsheet contains invalid row coordinates.")
            self.current_row = int(row)
            self.sheet_rows = max(self.sheet_rows, self.current_row)
            self.current_column = 0
            _enforce_shape_counts(self.sheet_columns, max(0, self.sheet_rows - 1))
        elif tag == "c" and "worksheet" in self.tags:
            if "r" in attributes:
                self._coordinate(attributes["r"])
            else:
                self.current_column += 1
                self.sheet_columns = max(self.sheet_columns, self.current_column)
                _enforce_shape_counts(self.sheet_columns, max(0, self.sheet_rows - 1))
            self.shared_cell = attributes.get("t") == "s"
            self.reference = ""
        elif tag == "si" and "sst" in self.tags:
            self.shared_size = 2

    def text(self, value: str) -> None:
        size = _json_string_size(value) - 2
        self.text_bytes += size
        if self.text_bytes > MAX_DATASET_NORMALIZED_BYTES:
            raise _limit("Spreadsheet exceeds the normalized byte budget.")
        if self.tags[-1:] == ["t"] and "si" in self.tags and "sst" in self.tags:
            self.shared_size += size
        if self.tags[-1:] == ["v"] and self.shared_cell:
            if len(self.reference) + len(value) > 16:
                raise DatasetParseError("Spreadsheet contains an invalid shared-string reference.")
            self.reference += value

    def end(self, name: str) -> None:
        tag = name.rsplit("}", 1)[-1]
        if tag == "si" and "sst" in self.tags:
            self.string_sizes.append(self.shared_size)
        elif tag == "v" and self.shared_cell:
            reference = self.reference.strip()
            if not reference.isascii() or not reference.isdigit():
                raise DatasetParseError("Spreadsheet contains an invalid shared-string reference.")
            index = int(reference)
            self.references[index] = self.references.get(index, 0) + 1
        elif tag == "c":
            self.shared_cell = False
        self.tags.pop()

    def reject_declaration(self, *args: Any) -> None:
        raise DatasetParseError("Spreadsheet XML entity or document declarations are unsafe.")

    def xml_parser(self):
        self.tags.clear()
        self.shared_cell = False
        xml_parser = expat.ParserCreate(namespace_separator="}")
        xml_parser.StartElementHandler = self.start
        xml_parser.EndElementHandler = self.end
        xml_parser.CharacterDataHandler = self.text
        xml_parser.StartDoctypeDeclHandler = self.reject_declaration
        xml_parser.EntityDeclHandler = self.reject_declaration
        xml_parser.ExternalEntityRefHandler = self.reject_declaration
        return xml_parser

    def finish(self) -> None:
        expanded = 0
        for index, count in self.references.items():
            if index >= len(self.string_sizes):
                raise DatasetParseError("Spreadsheet contains an invalid shared-string reference.")
            expanded += self.string_sizes[index] * count
            if expanded > MAX_DATASET_NORMALIZED_BYTES:
                raise _limit("Spreadsheet exceeds the normalized byte budget.")


def _preflight_zip_directory(content: bytes) -> None:
    end_offset = content.rfind(b"PK\x05\x06", max(0, len(content) - 65535 - _ZIP_END.size))
    if end_offset < 0 or end_offset + _ZIP_END.size > len(content):
        raise DatasetParseError("Spreadsheet archive directory is invalid.")
    end_record = _ZIP_END.unpack_from(content, end_offset)
    if end_record[1] or end_record[2] or end_record[3] != end_record[4]:
        raise DatasetParseError("Multi-volume spreadsheet archives are not supported.")
    if end_record[4] > MAX_XLSX_ARCHIVE_MEMBERS:
        raise _limit("Spreadsheet exceeds the archive member limit.")
    cursor = end_offset - end_record[5]
    if cursor < 0:
        raise DatasetParseError("Spreadsheet archive directory is invalid.")
    member_count = 0
    while cursor < end_offset:
        if cursor + _ZIP_MEMBER.size > end_offset:
            raise DatasetParseError("Spreadsheet archive directory is invalid.")
        member = _ZIP_MEMBER.unpack_from(content, cursor)
        if member[0] != b"PK\x01\x02":
            raise DatasetParseError("Spreadsheet archive directory is invalid.")
        member_count += 1
        if member_count > MAX_XLSX_ARCHIVE_MEMBERS:
            raise _limit("Spreadsheet exceeds the archive member limit.")
        cursor += _ZIP_MEMBER.size + member[10] + member[11] + member[12]
    if cursor != end_offset or member_count != end_record[4]:
        raise DatasetParseError("Spreadsheet archive directory is invalid.")


def _assert_xlsx_is_not_a_zip_bomb(content: bytes) -> None:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise DatasetParseError("Unsupported spreadsheet format. Use XLSX or CSV.")
    _preflight_zip_directory(content)
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_XLSX_ARCHIVE_MEMBERS:
                raise _limit("Spreadsheet exceeds the archive member limit.")
            if len({member.filename for member in members}) != len(members):
                raise DatasetParseError("Spreadsheet archive members must be unique.")
            uncompressed = sum(member.file_size for member in members)
            if uncompressed > MAX_XLSX_UNCOMPRESSED_BYTES or uncompressed > len(content) * MAX_XLSX_COMPRESSION_RATIO:
                raise _limit("Spreadsheet expands far beyond its file size and was rejected as unsafe.")
            preflight = _SpreadsheetPreflight()
            received = 0
            for member in members:
                if member.flag_bits & 1:
                    raise DatasetParseError("Encrypted spreadsheet archives are not supported.")
                if member.compress_type not in (zipfile.ZIP_STORED, zipfile.ZIP_DEFLATED, zipfile.ZIP_BZIP2, zipfile.ZIP_LZMA):
                    raise DatasetParseError("Spreadsheet archive compression is not supported.")
                xml_parser = preflight.xml_parser() if member.filename.endswith((".xml", ".rels")) else None
                member_received = 0
                with archive.open(member) as stream:
                    while chunk := stream.read(64 * 1024):
                        received += len(chunk)
                        member_received += len(chunk)
                        if received > MAX_XLSX_UNCOMPRESSED_BYTES:
                            raise _limit("Spreadsheet expands beyond the uncompressed byte limit.")
                        if xml_parser is not None:
                            xml_parser.Parse(chunk, False)
                            if member_received - xml_parser.CurrentByteIndex > MAX_XLSX_XML_TOKEN_BYTES:
                                raise _limit("Spreadsheet XML token exceeds the complexity budget.")
                    if xml_parser is not None:
                        xml_parser.Parse(b"", True)
            preflight.finish()
    except (zipfile.BadZipFile, expat.ExpatError, zlib.error, lzma.LZMAError, EOFError, OSError):
        raise DatasetParseError("Spreadsheet archive or XML is invalid.") from None


def _load_xlsx_workbook(content: bytes) -> Any:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError:
        raise DatasetParseError(
            "XLSX parsing requires the optional openpyxl reader; upload CSV or JSON instead.",
            error_code="dataset_format_unsupported", status_code=415,
        ) from None
    try:
        return load_workbook(io.BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except (InvalidFileException, zipfile.BadZipFile, ParseError, ValueError, KeyError, OSError):
        raise DatasetParseError("Spreadsheet workbook is invalid.") from None


def _parse_xlsx(content: bytes) -> tuple[list[str], list[dict[str, Any]]]:
    _assert_xlsx_is_not_a_zip_bomb(content)
    workbook = _load_xlsx_workbook(content)
    try:
        if not workbook.worksheets:
            raise DatasetParseError("Spreadsheet contains no worksheets.")
        sheet = workbook.worksheets[0]
        sheet.reset_dimensions()
        sheet.calculate_dimension(force=True)
        _enforce_shape_counts(sheet.max_column or 0, max(0, (sheet.max_row or 1) - 1))
        values = sheet.iter_rows(values_only=True)
        header = next(values, ())
        columns = _normalize_columns(["" if value is None else str(value) for value in header])
        budget = _NormalizedBudget(columns, max(0, (sheet.max_row or 1) - 1))
        rows: list[dict[str, Any]] = []
        for row in values:
            _enforce_shape_counts(len(columns), len(rows) + 1)
            record = dict(zip(columns, row, strict=True))
            budget.add_record(record)
            rows.append({column: _cast_value(value) for column, value in record.items()})
        if not rows:
            raise DatasetParseError("Dataset contains headers but 0 data rows.")
        return columns, rows
    finally:
        workbook.close()


def _cast_value(val: Any) -> Any:
    if val is None:
        return None
    if isinstance(val, float) and not math.isfinite(val):
        raise DatasetParseError("Dataset numeric values must be finite.")
    if isinstance(val, (int, float, bool, list, dict)):
        return val
    s = str(val).strip()
    if s == "" or s.lower() in ("null", "none", "nan", "na", "n/a"):
        return None
    if s.lower() == "true":
        return True
    if s.lower() == "false":
        return False
    # Try integer
    try:
        if s.isdigit() or (s.startswith("-") and s[1:].isdigit()):
            return int(s)
    except ValueError:
        pass
    # Try float
    try:
        numeric_value = float(s)
    except ValueError:
        return s
    if not math.isfinite(numeric_value):
        raise DatasetParseError("Dataset numeric values must be finite.")
    return numeric_value
